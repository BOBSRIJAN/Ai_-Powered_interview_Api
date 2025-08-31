from docx import Document
import PyPDF2

class CvToSimpleText():
    def extractTextFromDocx(filePath : str|None) -> str:
        doc = Document(filePath)
        textParts = []
        seen = set()
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if text and text not in seen:
                textParts.append(text)
                seen.add(text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and text not in seen:
                        textParts.append(text)
                        seen.add(text)
        return "\n".join(textParts)

    def extractTextFromPdf(filePath : str|None) -> str:
        textParts = []
        seen = set()

        with open(filePath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                pageText = page.extract_text()
                if pageText:
                    for line in pageText.splitlines():
                        cleanLine = line.strip()
                        if cleanLine and cleanLine not in seen:
                            textParts.append(cleanLine)
                            seen.add(cleanLine)
        return "\n".join(textParts)
    
# if __name__ == "__main__": # remove this in Production!
#     # Example using pdf 
#     pdf_text = CvToSimpleText.extractTextFromPdf(r"demoData\srijanraycv.pdf")
#     print(pdf_text)
    
#     # Example using docx
#     docx_text = CvToSimpleText.extractTextFromDocx(r"demoData\srijanraycv.docx")
#     print(docx_text)